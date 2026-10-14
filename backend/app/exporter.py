"""CSV / Excel / JSON exports."""
import csv, io, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .rules import SCHEMAS


def flat(doc):
    row = {"id": doc["id"], "file": doc["filename"], "doc_type": doc["doc_type"], "status": doc["status"]}
    for k, v in doc["data"]["fields"].items():
        row[k] = v.get("value")
    row["items"] = len(doc["data"].get("items", []))
    row["validation"] = "pass" if all(c["status"] == "pass" for c in doc["validation"] or []) else "issues"
    return row


def to_json(docs):
    return json.dumps([{"id": d["id"], "file": d["filename"], "doc_type": d["doc_type"], "status": d["status"], "approved_at": d["approved_at"],
                        "fields": {k: v.get("value") for k, v in d["data"]["fields"].items()},
                        "confidence": {k: v.get("confidence") for k, v in d["data"]["fields"].items()},
                        "items": d["data"].get("items", []), "validation": d["validation"]} for d in docs], indent=2, ensure_ascii=False)


def to_csv(docs):
    rows = [flat(d) for d in docs]
    cols = list(dict.fromkeys(k for r in rows for k in r)) or ["id"]
    buf = io.StringIO(); w = csv.DictWriter(buf, fieldnames=cols); w.writeheader(); w.writerows(rows)
    return buf.getvalue()


def to_xlsx(docs) -> bytes:
    wb = Workbook(); wb.remove(wb.active)
    head = PatternFill("solid", fgColor="D4AF37")
    for t, title, items_title in (("invoice", "Invoices", "Invoice lines"), ("receipt", "Receipts", "Receipt lines"), ("resume", "Resumes", "Experience")):
        ds = [d for d in docs if d["doc_type"] == t]
        if not ds:
            continue
        s = SCHEMAS[t]
        ws = wb.create_sheet(title); ws.append(["id", "file", "status"] + s["fields"] + ["validation"])
        for d in ds:
            f = flat(d); ws.append([f["id"], f["file"], f["status"]] + [f.get(k) for k in s["fields"]] + [f["validation"]])
        wi = wb.create_sheet(items_title); wi.append(["document_id", "file"] + s["items"])
        for d in ds:
            for it in d["data"].get("items", []):
                wi.append([d["id"], d["filename"]] + [it.get(k) for k in s["items"]])
        for sh in (ws, wi):
            for c in sh[1]:
                c.font = Font(bold=True); c.fill = head
            for col in sh.columns:
                sh.column_dimensions[col[0].column_letter].width = min(45, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    if not wb.sheetnames:
        wb.create_sheet("Empty").append(["No documents"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
