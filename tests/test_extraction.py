import io, json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from app import rules, text, validate

S = Path(__file__).resolve().parents[1] / "samples"


def ex(name):
    p = S / name
    return rules.extract(text.extract_text(p, ""), name)


def val(d, k):
    return d["fields"][k]["value"]


def test_invoice_inr():
    d = ex("invoice_nimbus_cloud.pdf")
    assert d["doc_type"] == "invoice"
    assert val(d, "vendor") == "Nimbus Cloud Services Pvt Ltd" and val(d, "invoice_number") == "NCS-2026-0412"
    assert val(d, "invoice_date") == "2026-09-03" and val(d, "currency") == "INR"
    assert (val(d, "subtotal"), val(d, "tax"), val(d, "total")) == (32300.0, 5814.0, 38114.0)
    assert len(d["items"]) == 3 and d["items"][1]["quantity"] == 2
    assert all(c["status"] == "pass" for c in validate.check(d))
    assert all(0 < f["confidence"] <= 1 for f in d["fields"].values())


def test_invoice_usd_long_dates():
    d = ex("invoice_bluepeak_design.pdf")
    assert val(d, "invoice_date") == "2026-09-21" and val(d, "due_date") == "2026-10-06" and val(d, "currency") == "USD"
    assert val(d, "total") == 4968.68


def test_validation_catches_bad_total():
    d = ex("invoice_harbor_supplies_mismatch.pdf")
    failed = {c["rule"] for c in validate.check(d) if c["status"] == "error"}
    assert failed == {"total_matches"}


def test_receipt_pdf():
    d = ex("receipt_cafe_aroma.pdf")
    assert d["doc_type"] == "receipt" and val(d, "merchant") == "CAFE AROMA" and val(d, "total") == 981.75
    assert val(d, "payment_method") == "UPI" and len(d["items"]) == 4


def test_receipt_image_ocr():
    pytest.importorskip("rapidocr_onnxruntime")
    d = ex("receipt_freshmart.png")
    assert d["doc_type"] == "receipt" and val(d, "total") == 1235.0 and val(d, "date") == "2026-10-08" and len(d["items"]) == 4


def test_resume():
    d = ex("resume_daniel_okafor.pdf")
    assert d["doc_type"] == "resume" and val(d, "name") == "Daniel Okafor" and val(d, "email") == "daniel.okafor@example.com"
    assert "FastAPI" in val(d, "skills") and len(d["items"]) == 3 and d["items"][0]["end"] == "Present"
    assert all(c["status"] == "pass" for c in validate.check(d))


def test_date_validation():
    d = {"doc_type": "invoice", "fields": {"invoice_date": {"value": "31/02/2026"}, "total": {"value": 1}}, "items": []}
    assert any(c["rule"] == "invoice_date_valid" and c["status"] == "error" for c in validate.check(d))


def test_batch_upload_review_approve_export(client):
    files = [("files", (n, (S / n).read_bytes(), "application/pdf")) for n in ("invoice_harbor_supplies_mismatch.pdf", "resume_ananya_iyer.pdf")]
    files.append(("files", ("notes.txt", b"hello", "text/plain")))
    r = client.post("/api/documents", files=files).json()
    inv, res, bad = r["documents"]
    assert bad["error"].startswith("Unsupported")
    assert client.post(f"/api/documents/{inv['id']}/approve").status_code == 422  # blocked by failing check
    fixed = client.put(f"/api/documents/{inv['id']}", json={"fields": {"total": 19809.84}}).json()
    assert fixed["data"]["fields"]["total"]["edited"] and all(c["status"] != "error" for c in fixed["validation"])
    assert client.post(f"/api/documents/{inv['id']}/approve").json()["status"] == "approved"
    assert client.post(f"/api/documents/{res['id']}/approve").json()["status"] == "approved"
    j = client.get("/api/export?format=json").json()
    assert {d["doc_type"] for d in j} == {"invoice", "resume"} and j[0]["fields"]["total"] == 19809.84
    csv = client.get("/api/export?format=csv").text
    assert "invoice_number" in csv.splitlines()[0] and "HOS/26/0931" in csv
    wb = load_workbook(io.BytesIO(client.get("/api/export?format=xlsx").content))
    assert {"Invoices", "Invoice lines", "Resumes", "Experience"} <= set(wb.sheetnames)
    assert wb["Invoice lines"].max_row == 4


def test_samples_endpoint(client):
    docs = client.post("/api/samples/load").json()["documents"]
    assert len(docs) == 7 and {d["doc_type"] for d in docs} == {"invoice", "receipt", "resume"}
    assert client.get("/api/health").json()["mode"] == "mock"
